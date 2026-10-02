"""
Weekly supervisor reports: Excel -> Word

Creates ONE Word document per supervisor. Each week (each sheet in the
workbook, e.g. "Aug 10 to 16") becomes ONE page with:
    - supervisor name, centered
    - the week
    - a table: Service level | Activity title | JIRA ID | Total hours
      grouped into Basic / Intermediate / Senior

Engineer levels come from a small Excel file you maintain by hand
(engineer_levels.xlsx with columns: Name | Level).

Usage:
    pip install pandas openpyxl python-docx
    python make_reports.py hours.xlsx
    python make_reports.py hours.xlsx --levels engineer_levels.xlsx --out reports
"""

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

LEVELS = ["Basic", "Intermediate", "Senior"]
HEADERS = ["Service level", "Activity title", "JIRA ID", "Total hours"]
COL_WIDTHS = [Cm(3.5), Cm(7.5), Cm(3.5), Cm(2.5)]

# Columns we need from each weekly sheet
COL_NAME = "Name"
COL_SUPERVISOR = "Supervisor"
COL_ACTIVITY = "Activity title"
COL_JIRA = "JIRA ID"


# ----------------------------------------------------------------- helpers
def norm(text):
    """Normalise a name for matching: trim, collapse spaces, ignore case."""
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def clean(value):
    """Turn Excel values into display text (blank for empty cells)."""
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def load_levels(path):
    df = pd.read_excel(path)
    df.columns = [str(c).strip() for c in df.columns]
    if "Name" not in df.columns or "Level" not in df.columns:
        sys.exit(f"'{path}' must have the columns 'Name' and 'Level'.")
    levels = {}
    for _, row in df.dropna(subset=["Name"]).iterrows():
        level = clean(row["Level"]).capitalize()
        if not level:                 # not filled in yet - reported as "no level" later
            continue
        if level not in LEVELS:
            print(f"  ! Unknown level '{row['Level']}' for {row['Name']} "
                  f"(use Basic, Intermediate or Senior) - ignored")
            continue
        levels[norm(row["Name"])] = level
    return levels


def find_columns(header_cells):
    """
    Work out which column is which, tolerating small wording differences
    (e.g. "Supervisor Name", "Engineer Name", "JIRA", "Activity Title ").
    Returns {standard name: column index} for the columns it found.
    """
    found = {}
    cells = [norm(c) for c in header_cells]

    def first(match):
        for idx, c in enumerate(cells):
            if c and match(c):
                return idx
        return None

    found[COL_SUPERVISOR] = first(lambda c: "supervisor" in c)
    found[COL_JIRA] = first(lambda c: "jira" in c)
    found[COL_ACTIVITY] = first(lambda c: "activity title" in c)
    if found[COL_ACTIVITY] is None:
        found[COL_ACTIVITY] = first(lambda c: "activity" in c
                                    and "description" not in c
                                    and "feature" not in c)
    # Engineer name: an exact "Name" column first, otherwise e.g. "Engineer Name"
    found[COL_NAME] = first(lambda c: c == "name")
    if found[COL_NAME] is None:
        others = ("supervisor", "program", "feature", "activity", "model")
        found[COL_NAME] = first(lambda c: "name" in c
                                and not any(o in c for o in others))
    return {k: v for k, v in found.items() if v is not None}


def read_week_sheet(xls, sheet):
    """
    Read a weekly sheet, finding the header row wherever it is.
    Returns (dataframe with standard column names, None) or (None, reason).
    """
    # keep_default_na=False: keep text like "N/A" or "NA" as it is written,
    # instead of pandas treating it as an empty cell
    raw = pd.read_excel(xls, sheet_name=sheet, header=None,
                        keep_default_na=False, na_values=[""])
    needed = [COL_NAME, COL_SUPERVISOR, COL_ACTIVITY, COL_JIRA]
    best = {}
    for i in range(min(60, len(raw))):
        cols = find_columns(raw.iloc[i].tolist())
        if len(cols) > len(best):
            best = cols
        if all(k in cols for k in needed):
            df = raw.iloc[i + 1:, [cols[k] for k in needed]].copy()
            df.columns = needed
            return df.dropna(how="all"), None

    # Couldn't find it: explain what the sheet looks like so it can be fixed
    lines = []
    if best:
        missing = [k for k in needed if k not in best]
        lines.append(f"closest header row is missing: {missing}")
    preview = raw.dropna(how="all").head(4)
    for _, row in preview.iterrows():
        vals = [clean(v) for v in row.tolist() if clean(v)]
        lines.append("first rows look like: " + " | ".join(vals)[:150])
    if raw.dropna(how="all").empty:
        lines.append("the sheet appears to be empty")
    return None, "\n      ".join(lines)


def set_cell_shading(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def write_cell(cell, text, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def write_bullets(cell, lines):
    """Write each line as a bullet point inside a single table cell."""
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for i, text in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.style = "List Bullet"
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text or "-")   # keep the bullets of both columns lined up
        run.font.size = Pt(10)


# -------------------------------------------------------------- the page
def add_week_page(doc, supervisor, week, activities_by_level, first_page):
    if not first_page:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(supervisor)
    r.bold = True
    r.font.size = Pt(16)

    wk = doc.add_paragraph()
    wk.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = wk.add_run(f"Week: {week}")
    r.font.size = Pt(12)

    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(HEADERS):
        cell = table.rows[0].cells[i]
        write_cell(cell, h, bold=True, center=True)
        set_cell_shading(cell, "D9E2F3")

    # One row per level; activities and JIRA IDs as bullet lists in one cell
    for level in LEVELS:
        items = activities_by_level.get(level, [])
        cells = table.add_row().cells
        write_cell(cells[0], level, bold=True, center=True)
        write_bullets(cells[1], [a for a, _ in items])
        write_bullets(cells[2], [j for _, j in items])
        write_cell(cells[3], "")              # Total hours: left blank for now

    table.autofit = False
    for i, w in enumerate(COL_WIDTHS):
        table.columns[i].width = w
    for row in table.rows:
        for i, w in enumerate(COL_WIDTHS):
            row.cells[i].width = w


# ------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Build weekly supervisor reports.")
    ap.add_argument("hours", help="Excel workbook with one sheet per week")
    ap.add_argument("--levels", default="engineer_levels.xlsx",
                    help="Excel file with columns Name | Level")
    ap.add_argument("--out", default="reports", help="Output folder")
    args = ap.parse_args()

    levels = load_levels(args.levels)
    xls = pd.ExcelFile(args.hours)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    # data[supervisor][week][level] = [(activity, jira), ...]  (unique, in order)
    data, week_order = {}, []
    unassigned = set()

    for sheet in xls.sheet_names:
        df, problem = read_week_sheet(xls, sheet)
        if df is None:
            print(f"  - Skipping sheet '{sheet}':\n      {problem}")
            continue
        week_order.append(sheet)

        for _, row in df.iterrows():
            supervisor = clean(row[COL_SUPERVISOR])
            name = clean(row[COL_NAME])
            if not supervisor or not name:
                continue
            level = levels.get(norm(name))
            if level is None:
                unassigned.add(name)
                continue
            entry = (clean(row[COL_ACTIVITY]), clean(row[COL_JIRA]))
            if not any(entry):
                continue
            bucket = (data.setdefault(supervisor, {})
                          .setdefault(sheet, {})
                          .setdefault(level, []))
            if entry not in bucket:          # same activity on several days -> once
                bucket.append(entry)

    if not data:
        sys.exit("No rows found. Check the sheet layout and the levels file.")

    for supervisor, weeks in data.items():
        doc = Document()
        for section in doc.sections:
            section.top_margin = section.bottom_margin = Cm(2)
            section.left_margin = section.right_margin = Cm(2)
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        first = True
        for week in week_order:
            if week in weeks:
                add_week_page(doc, supervisor, week, weeks[week], first)
                first = False

        safe = re.sub(r'[\\/:*?"<>|]+', "_", supervisor)
        path = out_dir / f"Report - {safe}.docx"
        doc.save(path)
        print(f"  ✓ {path}  ({len(weeks)} week(s))")

    if unassigned:
        print("\n  ! These engineers have no level in the levels file and were left out:")
        for n in sorted(unassigned):
            print(f"      - {n}")
        print("    Add them to the levels file and run again.")


if __name__ == "__main__":
    main()