"""
Monthly supervisor reports: Excel -> Word

Same idea as make_reports.py, for the monthly workbook layout:
    Date | Name | CSID | Ford TL | Ford Supervisor | Feature name |
    JIRA ID/Test Rail Test Run | Activity Description |
    TC executed, TC created/Updated, ... | Program | (anything else is ignored)

Creates ONE Word document per supervisor with:
    - the supervisor name in a shaded box at the top
    - for every week: "Week 39: 01/09/2026 - 06/09/2026  - Supervisor: <name>"
      and a table with one row each for Basic / Intermediate / Senior

Excel column               -> Word column
    Activity Description       -> Activity Title
    JIRA ID/Test Rail Test Run -> JIRA ID/TMT Request ID/Test Rail Test Run
    TC executed, ...           -> Activity Description
    TC executed, ...           -> TC executed, ...
    Program                    -> Vehicle Program
    (none)                     -> Total Amount of hours (left blank for now)

A JIRA cell with several tickets ("ABC-1, ABC-2 and ABC-3") becomes one bullet
per ticket, repeating the activity title and the other columns. If the title
cell has as many titles as there are tickets, they are paired in order.

Weeks come from the Date column, not from the sheet names: Monday to Sunday,
cut at the start and end of the month (so September 2026 gives 01/09 - 06/09,
07/09 - 13/09, ..., 28/09 - 30/09). Week numbers are ISO week numbers unless
--first-week is given, which numbers the first week N and counts up from there.
A sheet without a Date column is treated as one week named after the sheet,
like in make_reports.py.

Engineer levels come from the same engineer_levels.xlsx as make_reports.py
(make_levels.py also works on the monthly workbook).

Usage:
    pip install pandas openpyxl python-docx
    python make_monthly_reports.py monthly.xlsx
    python make_monthly_reports.py monthly.xlsx --levels engineer_levels.xlsx --out monthly_reports
    python make_monthly_reports.py monthly.xlsx --first-week 39
"""

import argparse
import calendar
import re
import sys
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

from make_reports import (LEVELS, COL_NAME, COL_SUPERVISOR, COL_JIRA,
                          norm, clean, load_levels, read_week_sheet,
                          set_cell_shading, write_cell)

HEADERS = ["Service Level", "Activity Title",
           "JIRA ID/TMT Request ID/Test Rail Test Run", "Activity Description",
           "TC executed, TC created/Updated, System Issues Retested/Reported, "
           "Test Procedure issues resolved, Deliverables completed",
           "Vehicle Program", "Total Amount of hours"]
COL_WIDTHS = [Cm(2.6), Cm(3.1), Cm(2.5), Cm(2.6), Cm(3.1), Cm(2.6), Cm(2.0)]
TITLE_FILL = "F6C5AC"
SPLIT_COLOR = RGBColor(0xC0, 0x00, 0x00)   # items split from a multi-ticket cell

# Columns we need from the monthly sheet (plus Name / Supervisor / JIRA ID)
COL_DATE = "Date"
COL_DESCRIPTION = "Activity Description"
COL_TC = "TC executed"
COL_PROGRAM = "Program"
# Excel column feeding each Word column after "Service Level", in order
# (None = left blank). Total Amount of hours is always blank.
ENTRY_COLS = [COL_DESCRIPTION,   # -> Activity Title
              COL_JIRA,          # -> JIRA ID/TMT Request ID/Test Rail Test Run
              COL_TC,            # -> Activity Description
              COL_TC,            # -> TC executed, ...
              COL_PROGRAM]       # -> Vehicle Program
# dict.fromkeys: each Excel column once, even if it feeds two Word columns
NEEDED = list(dict.fromkeys([COL_NAME, COL_SUPERVISOR] + [c for c in ENTRY_COLS if c]))


# ----------------------------------------------------------------- helpers
def find_columns(header_cells):
    """
    Work out which column is which in the monthly layout, tolerating small
    wording differences. Returns {standard name: column index}.
    """
    cells = [norm(c) for c in header_cells]

    def first(*matches):
        for match in matches:                  # try the strictest match first
            for idx, c in enumerate(cells):
                if c and match(c):
                    return idx
        return None

    others = ("supervisor", "program", "feature", "activity", "model")
    found = {
        # \b so "TC created/Updated" doesn't count as a date column
        COL_DATE: first(lambda c: c == "date", lambda c: re.search(r"\bdate\b", c)),
        COL_NAME: first(lambda c: c == "name",
                        lambda c: "name" in c and not any(o in c for o in others)),
        COL_SUPERVISOR: first(lambda c: "supervisor" in c),
        COL_JIRA: first(lambda c: "jira" in c),
        COL_DESCRIPTION: first(lambda c: "description" in c),
        COL_TC: first(lambda c: c.startswith("tc executed"),
                      lambda c: "tc executed" in c),
        COL_PROGRAM: first(lambda c: c == "program", lambda c: "program" in c),
    }
    return {k: v for k, v in found.items() if v is not None}


def to_date(value):
    """Excel date cell (or text / serial number) -> date, or None."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if pd.isna(value):
            return None
        value = pd.Timestamp("1899-12-30") + pd.Timedelta(days=value)
    d = pd.to_datetime(value, errors="coerce")
    return None if pd.isna(d) else d.date()


TICKET_RE = re.compile(r"\b[A-Za-z][A-Za-z0-9_]*-\d+\b")       # ABC-123
TICKET_SEP = re.compile(r"[\s,;&]+")
TITLE_SEP = re.compile(r"\s*(?:[,;\n•&]|\s+and\s+|\s+y\s+)\s*", re.I)


def split_tickets(text):
    """
    "ABC-1, ABC-2 and ABC-3" -> ["ABC-1", "ABC-2", "ABC-3"].
    A single ticket, "N/A" or free text comes back as [text].
    """
    keys = list(dict.fromkeys(TICKET_RE.findall(text)))
    if len(keys) > 1:
        return keys
    # Test Rail runs / plain numbers: "R1234 R1235", "12345, 12346"
    parts = [p for p in TICKET_SEP.split(text) if p and p.lower() not in ("and", "y")]
    if len(parts) > 1 and all(re.search(r"\d", p) for p in parts):
        return list(dict.fromkeys(parts))
    return [text]


def split_entry(entry):
    """
    One Excel row -> one entry per ticket in its JIRA cell, the other columns
    repeated. Several titles in the title cell are paired with the tickets in
    order when the counts match; otherwise every ticket gets the whole title.
    """
    jira_i = ENTRY_COLS.index(COL_JIRA)
    title_i = ENTRY_COLS.index(COL_DESCRIPTION)
    tickets = split_tickets(entry[jira_i])
    if len(tickets) == 1:
        return [entry]
    titles = [t for t in TITLE_SEP.split(entry[title_i]) if t]
    if len(titles) != len(tickets):
        titles = [entry[title_i]] * len(tickets)
    parts = []
    for title, ticket in zip(titles, tickets):
        part = list(entry)
        part[title_i], part[jira_i] = title, ticket
        parts.append(tuple(part))
    return parts


def week_label(number, monday, year, month):
    """"Week 39: 01/09/2026 - 06/09/2026" - Monday to Sunday, cut to the month."""
    start = max(monday, date(year, month, 1))
    end = min(monday + timedelta(days=6),
              date(year, month, calendar.monthrange(year, month)[1]))
    return f"Week {number}: {start:%d/%m/%Y} - {end:%d/%m/%Y}"


# -------------------------------------------------------------- the page
def add_title_box(doc, supervisor):
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.rows[0].cells[0]
    cell.width = Cm(13.75)
    set_cell_shading(cell, TITLE_FILL)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(supervisor)
    r.font.name = "Calibri Light"
    r.font.size = Pt(16)
    doc.add_paragraph()


def repeat_as_header(row):
    """Repeat this row at the top of every page the table runs onto."""
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def write_numbered(cell, lines, flagged):
    """
    Write the lines as 1., 2., 3. ... inside one table cell (numbering starts
    at 1 in every cell, so item N lines up across the columns). Lines whose
    flag is set (split from a cell with several tickets) are red and bold.
    """
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for i, (text, flag) in enumerate(zip(lines, flagged)):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(f"{i + 1}.\t{text or '-'}")
        run.font.size = Pt(10)
        if flag:
            run.bold = True
            run.font.color.rgb = SPLIT_COLOR


def add_week(doc, supervisor, label, activities_by_level):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{label}  - ")
    p.add_run(f"Supervisor: {supervisor}").bold = True

    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(HEADERS):
        write_cell(table.rows[0].cells[i], h)
    repeat_as_header(table.rows[0])

    # One row per level; each column a numbered list, one item per activity
    for level in LEVELS:
        items = activities_by_level.get(level, {})
        flagged = list(items.values())
        cells = table.add_row().cells
        write_cell(cells[0], level, bold=True, center=True)
        for col, source in enumerate(ENTRY_COLS):
            if source:
                write_numbered(cells[col + 1], [item[col] for item in items], flagged)
            else:
                write_cell(cells[col + 1], "")
        write_cell(cells[-1], "")             # Total hours: left blank for now

    table.autofit = False
    for i, w in enumerate(COL_WIDTHS):
        table.columns[i].width = w
    for row in table.rows:
        for i, w in enumerate(COL_WIDTHS):
            row.cells[i].width = w

    doc.add_paragraph()


# ------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Build monthly supervisor reports.")
    ap.add_argument("hours", help="Monthly Excel workbook")
    ap.add_argument("--levels", default="engineer_levels.xlsx",
                    help="Excel file with columns Name | Level")
    ap.add_argument("--out", default="monthly_reports", help="Output folder")
    ap.add_argument("--first-week", type=int,
                    help="Number of the first week (e.g. 39); the rest count up "
                         "from it. Default: ISO week numbers")
    args = ap.parse_args()

    levels = load_levels(args.levels)
    xls = pd.ExcelFile(args.hours)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    # data[supervisor][week key][level] = {entry following ENTRY_COLS: split?, ...}
    # (a dict keeps the entries unique and in order; split? = came from a cell
    # with several tickets, shown in red)
    # week key: (0, monday ordinal, year, month) for dated rows - a week that
    # runs into the next month is split in two - or (1, sheet index) for undated sheets
    data, labels = {}, {}
    unassigned = set()
    bad_dates = 0

    for sheet_idx, sheet in enumerate(xls.sheet_names):
        df, problem = read_week_sheet(xls, sheet, needed=NEEDED,
                                      optional=[COL_DATE], finder=find_columns)
        if df is None:
            print(f"  - Skipping sheet '{sheet}':\n      {problem}")
            continue
        dated = COL_DATE in df.columns

        for _, row in df.iterrows():
            supervisor = clean(row[COL_SUPERVISOR])
            name = clean(row[COL_NAME])
            if not supervisor or not name:
                continue
            level = levels.get(norm(name))
            if level is None:
                unassigned.add(name)
                continue
            entry = tuple(clean(row[c]) if c else "" for c in ENTRY_COLS)
            if not any(entry):
                continue

            if dated:
                d = to_date(row[COL_DATE])
                if d is None:
                    bad_dates += 1
                    continue
                monday = d - timedelta(days=d.weekday())
                week = (0, monday.toordinal(), d.year, d.month)
                labels[week] = monday
            else:
                week = (1, sheet_idx)
                labels[week] = sheet

            bucket = (data.setdefault(supervisor, {})
                          .setdefault(week, {})
                          .setdefault(level, {}))
            parts = split_entry(entry)       # several tickets in one cell -> one each
            for part in parts:               # same activity on several days -> once
                bucket[part] = bucket.get(part, False) or len(parts) > 1

    if not data:
        sys.exit("No rows found. Check the sheet layout and the levels file.")

    first_monday = min((w[1] for w in labels if w[0] == 0), default=0)
    for week, monday in labels.items():
        if week[0] == 0:
            if args.first_week is None:
                number = monday.isocalendar()[1]
            else:
                number = args.first_week + (week[1] - first_monday) // 7
            labels[week] = week_label(number, monday, week[2], week[3])

    for supervisor, weeks in data.items():
        doc = Document()
        for section in doc.sections:
            section.top_margin = section.bottom_margin = Cm(2.5)
            section.left_margin = section.right_margin = Cm(1.5)
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)

        add_title_box(doc, supervisor)
        for week in sorted(weeks):
            add_week(doc, supervisor, labels[week], weeks[week])

        safe = re.sub(r'[\\/:*?"<>|]+', "_", supervisor)
        path = out_dir / f"Monthly Report - {safe}.docx"
        try:
            doc.save(path)
        except PermissionError:
            sys.exit(f"Can't write '{path}' - close it in Word and run again.")
        print(f"  ✓ {path}  ({len(weeks)} week(s))")

    if bad_dates:
        print(f"\n  ! {bad_dates} row(s) had an empty or unreadable Date and were left out.")
    if unassigned:
        print("\n  ! These engineers have no level in the levels file and were left out:")
        for n in sorted(unassigned):
            print(f"      - {n}")
        print("    Add them to the levels file (python make_levels.py <workbook>) and run again.")


if __name__ == "__main__":
    main()
