"""
Shared engine for make_reports.py, make_monthly_reports.py, make_levels.py
and the window (report_gui.pyw).

Everything that can be changed in the window lives in a *profile* (a dict):

    layout          "weekly"  - one page per week: supervisor, week, table
                    "monthly" - supervisor in a shaded box, then every week
                                as a bullet line followed by its table
    weeks_from      "sheet" - each sheet is one week, named after the sheet
                    "date"  - weeks come from the Date column: Monday to
                              Sunday, cut at the start and end of the month.
                              A sheet without a Date column counts as one week.
    first_week      number of the first week (date weeks only); None = ISO
    hours_file      the Excel workbook with the hours
    levels_file     Excel file with columns Name | Level
    out_dir         output folder
    file_name       Word file name; "{supervisor}" is replaced by the name
    supervisor_mode "column" - the supervisor is read from its Excel column
                    "fixed"  - the column is ignored and every row belongs
                               to supervisor_name
    fields          Excel columns to look for, in order:
                    [{name, words, ignore, required}, ...]
                    A header matches a field if it is exactly one of `words`
                    or contains one of them (as the start of a word), and
                    contains none of `ignore`. Not case-sensitive. Name,
                    Supervisor and Date are built in.
    columns         The Word table, left to right: [{header, source, width}]
                    source is LEVEL_SOURCE (Basic / Intermediate / Senior),
                    BLANK_SOURCE (left empty) or the name of a field.
"""

import calendar
import copy
import json
import re
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

LEVELS = ["Basic", "Intermediate", "Senior"]

# Built-in Excel fields
NAME = "Name"
SUPERVISOR = "Supervisor"
DATE = "Date"
BUILT_IN = (NAME, SUPERVISOR, DATE)

# Word column sources that don't come from Excel
LEVEL_SOURCE = "@level"
BLANK_SOURCE = "@blank"

WEEKLY_HEADER_FILL = "D9E2F3"
MONTHLY_TITLE_FILL = "F6C5AC"

# Page width of the default Word template (Letter) and the margins per layout
PAGE_WIDTH_CM = 21.59
MARGINS_CM = {"weekly": (2, 2), "monthly": (2.5, 1.5)}   # (top/bottom, left/right)


class ReportError(Exception):
    """A problem the user can fix (missing file, bad layout, file open in Word...)."""


# ------------------------------------------------------------ profiles
def _field(name, words, ignore=(), required=True):
    return {"name": name, "words": list(words), "ignore": list(ignore),
            "required": required}


def _col(header, source, width):
    return {"header": header, "source": source, "width": width}


# Engineer name: "Name", or e.g. "Engineer Name" - but not "Supervisor Name",
# "Feature name", ...
_NAME_IGNORE = ["supervisor", "program", "feature", "activity", "model"]

DEFAULT_PROFILES = {
    "Weekly": {
        "layout": "weekly",
        "weeks_from": "sheet",
        "first_week": None,
        "hours_file": "",
        "levels_file": "engineer_levels.xlsx",
        "out_dir": "reports",
        "file_name": "Report - {supervisor}.docx",
        "supervisor_mode": "column",
        "supervisor_name": "",
        "fields": [
            _field(NAME, ["name"], _NAME_IGNORE),
            _field(SUPERVISOR, ["supervisor"]),
            _field(DATE, ["date"], required=False),
            _field("Activity title", ["activity title", "activity"],
                   ["description", "feature"]),
            _field("JIRA ID", ["jira"]),
        ],
        "columns": [
            _col("Service level", LEVEL_SOURCE, 3.5),
            _col("Activity title", "Activity title", 7.5),
            _col("JIRA ID", "JIRA ID", 3.5),
            _col("Total hours", BLANK_SOURCE, 2.5),
        ],
    },
    "Monthly": {
        "layout": "monthly",
        "weeks_from": "date",
        "first_week": None,
        "hours_file": "",
        "levels_file": "engineer_levels.xlsx",
        "out_dir": "monthly_reports",
        "file_name": "Monthly Report - {supervisor}.docx",
        "supervisor_mode": "column",
        "supervisor_name": "",
        "fields": [
            _field(NAME, ["name"], _NAME_IGNORE),
            _field(SUPERVISOR, ["supervisor"]),
            _field(DATE, ["date"], required=False),
            _field("Activity Description", ["description"]),
            _field("JIRA ID", ["jira"]),
            _field("TC executed", ["tc executed"]),
            _field("Program", ["program"]),
        ],
        "columns": [
            _col("Service Level", LEVEL_SOURCE, 2.6),
            _col("Activity Title", "Activity Description", 3.1),
            _col("JIRA ID/TMT Request ID/Test Rail Test Run", "JIRA ID", 2.5),
            _col("Activity Description", "TC executed", 2.6),
            _col("TC executed, TC created/Updated, System Issues Retested/Reported, "
                 "Test Procedure issues resolved, Deliverables completed",
                 "TC executed", 3.1),
            _col("Vehicle Program", "Program", 2.6),
            _col("Total Amount of hours", BLANK_SOURCE, 2.0),
        ],
    },
}


def default_profile(name_or_layout):
    """A fresh copy of the "Weekly" / "Monthly" defaults (also by layout name)."""
    key = {"weekly": "Weekly", "monthly": "Monthly"}.get(name_or_layout, name_or_layout)
    return copy.deepcopy(DEFAULT_PROFILES[key])


def complete_profile(profile):
    """Fill in anything missing (e.g. a settings file from an older version)."""
    base = default_profile(profile.get("layout", "weekly"))
    for key, value in base.items():
        profile.setdefault(key, value)
    names = [f["name"] for f in profile["fields"]]
    for i, f in enumerate(base["fields"]):
        if f["name"] in BUILT_IN and f["name"] not in names:
            profile["fields"].insert(i, f)
    return profile


def load_settings(path):
    """{"profile": current profile name, "profiles": {name: profile}}"""
    settings = {}
    try:
        settings = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        pass
    profiles = settings.get("profiles") or copy.deepcopy(DEFAULT_PROFILES)
    for p in profiles.values():
        complete_profile(p)
    current = settings.get("profile")
    if current not in profiles:
        current = next(iter(profiles))
    return {"profile": current, "profiles": profiles}


def save_settings(path, settings):
    Path(path).write_text(json.dumps(settings, indent=2, ensure_ascii=False),
                          encoding="utf-8")


def usable_width_cm(layout):
    return PAGE_WIDTH_CM - 2 * MARGINS_CM[layout][1]


# ------------------------------------------------------------ helpers
def norm(text):
    """Normalise a name for matching: trim, collapse spaces, ignore case."""
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def clean(value):
    """Turn Excel values into display text (blank for empty cells)."""
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _contains_word(header, word):
    # The word has to start a word in the header, so "date" finds "Date" and
    # "Work date" but not "TC created/Updated"
    return re.search(r"(?<!\w)" + re.escape(word), header) is not None


def active_fields(profile):
    """The fields to look for with these settings, in priority order."""
    out = []
    for f in profile["fields"]:
        if f["name"] == SUPERVISOR and profile["supervisor_mode"] == "fixed":
            continue
        if f["name"] == DATE and profile["weeks_from"] != "date":
            continue
        out.append(f)
    return out


def is_required(field):
    if field["name"] in (NAME, SUPERVISOR):
        return True
    if field["name"] == DATE:
        return False                   # without it the sheet counts as one week
    return bool(field.get("required", True))


def make_finder(fields):
    """
    Return finder(header cells) -> {field name: column index} for the fields
    it found. Fields are matched in order: exact header first, then a header
    containing one of the words. A column is only used for one field.
    """
    fields = [(f["name"],
               [norm(w) for w in f.get("words", []) if norm(w)],
               [norm(w) for w in f.get("ignore", []) if norm(w)])
              for f in fields]

    def finder(header_cells):
        cells = [norm(c) for c in header_cells]
        found, taken = {}, set()
        for name, words, ignore in fields:
            idx = None
            for w in words:                                  # exact match
                idx = next((i for i, c in enumerate(cells)
                            if c == w and i not in taken), None)
                if idx is not None:
                    break
            if idx is None:
                for w in words:                              # in order of priority
                    idx = next((i for i, c in enumerate(cells)
                                if c and i not in taken and _contains_word(c, w)
                                and not any(x in c for x in ignore)), None)
                    if idx is not None:
                        break
            if idx is not None:
                found[name] = idx
                taken.add(idx)
        return found

    return finder


def read_raw(xls, sheet):
    # keep_default_na=False: keep text like "N/A" or "NA" as it is written,
    # instead of pandas treating it as an empty cell
    return pd.read_excel(xls, sheet_name=sheet, header=None,
                         keep_default_na=False, na_values=[""])


def find_header(raw, needed, finder):
    """
    Look for the header row in the first 60 rows.
    Returns (row, {field: column index}, True) when every needed field is
    there, otherwise (closest row or None, what it found, False).
    """
    best, best_row = {}, None
    for i in range(min(60, len(raw))):
        cols = finder(raw.iloc[i].tolist())
        if len(cols) > len(best):
            best, best_row = cols, i
        if all(k in cols for k in needed):
            return i, cols, True
    return best_row, best, False


def preview_rows(raw, count=4):
    lines = []
    for _, row in raw.dropna(how="all").head(count).iterrows():
        vals = [clean(v) for v in row.tolist() if clean(v)]
        lines.append(" | ".join(vals)[:150])
    return lines


def read_week_sheet(xls, sheet, needed, optional=(), finder=None):
    """
    Read a sheet, finding the header row wherever it is.
    `needed` fields must all be found; `optional` ones are kept if present.
    Returns (dataframe with field names as columns, None) or (None, reason).
    """
    raw = read_raw(xls, sheet)
    needed = list(needed)
    row, cols, ok = find_header(raw, needed, finder)
    if ok:
        keep = needed + [k for k in optional if k in cols]
        df = raw.iloc[row + 1:, [cols[k] for k in keep]].copy()
        df.columns = keep
        return df.dropna(how="all"), None

    # Couldn't find it: explain what the sheet looks like so it can be fixed
    lines = []
    if cols:
        missing = [k for k in needed if k not in cols]
        lines.append(f"closest header row is missing: {missing}")
    lines += ["first rows look like: " + p for p in preview_rows(raw)]
    if raw.dropna(how="all").empty:
        lines.append("the sheet appears to be empty")
    return None, "\n      ".join(lines)


def open_workbook(path):
    if not str(path).strip():
        raise ReportError("Choose the Excel workbook with the hours first.")
    path = Path(path)
    if not path.exists():
        raise ReportError(f"Excel workbook not found: '{path}'")
    try:
        return pd.ExcelFile(path)
    except PermissionError:
        raise ReportError(f"Can't open '{path}' - close it in Excel and try again.")
    except Exception as e:
        raise ReportError(f"Can't read '{path}' as an Excel workbook: {e}")


# ------------------------------------------------------------- levels
def _read_levels_file(path):
    if not str(path).strip():
        raise ReportError("Choose the levels file (Name | Level) first.")
    path = Path(path)
    if not path.exists():
        raise ReportError(f"Levels file not found: '{path}'. "
                          "Create it with 'Create / update levels file' "
                          "(or python make_levels.py <workbook>).")
    try:
        df = pd.read_excel(path)
    except PermissionError:
        raise ReportError(f"Can't open '{path}' - close it in Excel and try again.")
    df.columns = [str(c).strip() for c in df.columns]
    if "Name" not in df.columns or "Level" not in df.columns:
        raise ReportError(f"'{path}' must have the columns 'Name' and 'Level'.")
    return df


def load_levels(path, log=print):
    levels = {}
    for _, row in _read_levels_file(path).dropna(subset=["Name"]).iterrows():
        level = clean(row["Level"]).capitalize()
        if not level:                 # not filled in yet - reported as "no level" later
            continue
        if level not in LEVELS:
            log(f"  ! Unknown level '{row['Level']}' for {row['Name']} "
                f"(use Basic, Intermediate or Senior) - ignored")
            continue
        levels[norm(row["Name"])] = level
    return levels


def load_roster(path):
    """Every engineer in the levels file, as written there: {norm(name): name}."""
    df = _read_levels_file(path)
    return {norm(n): clean(n) for n in df["Name"].dropna() if clean(n)}


def report_missing(weeks, uploaded, roster, log=print):
    """
    Log, week by week, the engineers who didn't upload any activity.
    weeks:    [(week key, label), ...] in order
    uploaded: {week key: {norm(name), ...}}
    roster:   {norm(name): name} - everyone expected to upload
    """
    missing = []
    for week, label in weeks:
        names = [roster[n] for n in roster if n not in uploaded.get(week, set())]
        if names:
            missing.append((label, sorted(names, key=str.lower)))
    if not missing:
        log("\n  ✓ Every engineer uploaded activities every week.")
        return
    log("\n  ! These engineers didn't upload any activities:")
    for label, names in missing:
        log(f"      {label}")
        for n in names:
            log(f"          - {n}")


# --------------------------------------------------------------- weeks
def to_date(value):
    """Excel date cell (or text / serial number) -> date, or None."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if pd.isna(value):
            return None
        value = pd.Timestamp("1899-12-30") + pd.Timedelta(days=value)
    d = pd.to_datetime(value, errors="coerce")
    return None if pd.isna(d) else d.date()


def week_label(number, monday, year, month):
    """"Week 39: 01/09/2026 - 06/09/2026" - Monday to Sunday, cut to the month."""
    start = max(monday, date(year, month, 1))
    end = min(monday + timedelta(days=6),
              date(year, month, calendar.monthrange(year, month)[1]))
    return f"Week {number}: {start:%d/%m/%Y} - {end:%d/%m/%Y}"


# ---------------------------------------------------------- Word helpers
def set_cell_shading(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def format_table_paragraph(p):
    """
    Table text: 10 pt, 0 pt before / after, single line spacing. The paragraph
    mark gets 10 pt as well, so empty cells and bullet symbols aren't 11 pt.
    """
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p_pr = p._p.get_or_add_pPr()
    r_pr = p_pr.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        p_pr.append(r_pr)
    for tag in ("w:sz", "w:szCs"):
        sz = r_pr.find(qn(tag))
        if sz is None:
            sz = OxmlElement(tag)
            r_pr.append(sz)
        sz.set(qn("w:val"), "20")             # half-points: 20 = 10 pt


def write_cell(cell, text, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    format_table_paragraph(p)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def write_bullets(cell, lines):
    """Write each line as a bullet point inside a single table cell."""
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    format_table_paragraph(cell.paragraphs[0])    # also when there are no lines
    for i, text in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.style = "List Bullet"
        format_table_paragraph(p)
        run = p.add_run(text or "-")   # keep the bullets of all columns lined up
        run.font.size = Pt(10)


def repeat_as_header(row):
    """Repeat this row at the top of every page the table runs onto."""
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


# ------------------------------------------------------------ the pages
def add_level_table(doc, columns, entry_index, activities_by_level, layout):
    """
    One row per level; each Excel-filled column is a bullet list with one
    bullet per activity. entry_index[i] = position in the activity tuple of
    the field feeding column i.
    """
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, col in enumerate(columns):
        cell = table.rows[0].cells[i]
        if layout == "weekly":
            write_cell(cell, col["header"], bold=True, center=True)
            set_cell_shading(cell, WEEKLY_HEADER_FILL)
        else:
            write_cell(cell, col["header"])
    if layout != "weekly":
        repeat_as_header(table.rows[0])

    for level in LEVELS:
        items = activities_by_level.get(level, [])
        cells = table.add_row().cells
        for i, col in enumerate(columns):
            if col["source"] == LEVEL_SOURCE:
                write_cell(cells[i], level, bold=True, center=True)
            elif col["source"] == BLANK_SOURCE:
                write_cell(cells[i], "")      # e.g. Total hours: filled in by hand
            else:
                write_bullets(cells[i], [item[entry_index[i]] for item in items])

    widths = [Cm(float(col["width"])) for col in columns]
    table.autofit = False
    for i, w in enumerate(widths):
        table.columns[i].width = w
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


def add_week_page(doc, supervisor, title, columns, entry_index,
                  activities_by_level, first_page):
    """Weekly layout: supervisor, week and the table on a page of their own."""
    if not first_page:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(supervisor)
    r.bold = True
    r.font.size = Pt(16)

    wk = doc.add_paragraph()
    wk.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = wk.add_run(title)
    r.font.size = Pt(12)

    add_level_table(doc, columns, entry_index, activities_by_level, "weekly")


def add_title_box(doc, supervisor):
    """Monthly layout: the supervisor name in a shaded box at the top."""
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.rows[0].cells[0]
    cell.width = Cm(13.75)
    set_cell_shading(cell, MONTHLY_TITLE_FILL)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(supervisor)
    r.font.name = "Calibri Light"
    r.font.size = Pt(16)
    doc.add_paragraph()


def add_week(doc, supervisor, label, columns, entry_index, activities_by_level):
    """Monthly layout: "<week>  - Supervisor: <name>" and the table."""
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{label}  - ")
    p.add_run(f"Supervisor: {supervisor}").bold = True
    add_level_table(doc, columns, entry_index, activities_by_level, "monthly")
    doc.add_paragraph()


def new_document(layout):
    doc = Document()
    top, side = MARGINS_CM[layout]
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(top)
        section.left_margin = section.right_margin = Cm(side)
    style = doc.styles["Normal"]
    style.font.name = "Calibri" if layout == "weekly" else "Times New Roman"
    style.font.size = Pt(11)
    return doc


def output_path(out_dir, template, supervisor, several):
    safe = re.sub(r'[\\/:*?"<>|]+', "_", supervisor)
    template = (template or "").strip() or "Report - {supervisor}.docx"
    if "{supervisor}" not in template and several:   # don't overwrite each other
        stem = template[:-5] if template.lower().endswith(".docx") else template
        template = stem + " - {supervisor}.docx"
    name = template.replace("{supervisor}", safe)
    name = re.sub(r'[\\/:*?"<>|]+', "_", name)
    if not name.lower().endswith(".docx"):
        name += ".docx"
    return Path(out_dir) / name


# ---------------------------------------------------------------- build
def check_profile(profile):
    """Raise ReportError for settings that can't work."""
    if not profile["columns"]:
        raise ReportError("The Word table has no columns - add some in 'Word table'.")
    names = {f["name"] for f in profile["fields"]}
    for col in profile["columns"]:
        src = col["source"]
        if src not in (LEVEL_SOURCE, BLANK_SOURCE) and src not in names:
            raise ReportError(f"Word column '{col['header']}' is filled from '{src}', "
                              "which isn't in the Excel columns any more.")
        try:
            if float(col["width"]) <= 0:
                raise ValueError
        except (TypeError, ValueError):
            raise ReportError(f"Word column '{col['header']}' needs a width above 0.")
    if profile["supervisor_mode"] == "fixed" and not profile["supervisor_name"].strip():
        raise ReportError("Write the supervisor's name, or read it from the "
                          "supervisor column instead.")
    for f in active_fields(profile):
        if not [w for w in f.get("words", []) if norm(w)]:
            raise ReportError(f"Excel column '{f['name']}' has no words to look for.")


def build_reports(profile, log=print):
    """Read the workbook and write one Word document per supervisor.
    Returns the paths written."""
    check_profile(profile)
    layout = profile["layout"]
    fixed = profile["supervisor_mode"] == "fixed"
    fixed_name = profile["supervisor_name"].strip()
    by_date = profile["weeks_from"] == "date"
    columns = profile["columns"]

    # Each activity is a tuple of the Excel fields shown in Word (each once,
    # even if it feeds two Word columns); entry_index maps column -> position
    entry_fields = list(dict.fromkeys(c["source"] for c in columns
                                      if c["source"] not in (LEVEL_SOURCE, BLANK_SOURCE)))
    entry_index = [entry_fields.index(c["source"]) if c["source"] in entry_fields else None
                   for c in columns]

    fields = active_fields(profile)
    needed = [f["name"] for f in fields if is_required(f)]
    optional = [f["name"] for f in fields if not is_required(f)]
    finder = make_finder(fields)

    levels = load_levels(profile["levels_file"], log)
    roster = load_roster(profile["levels_file"])   # + everyone found in the workbook
    xls = open_workbook(profile["hours_file"])
    out_dir = Path(profile["out_dir"] or ".")
    out_dir.mkdir(parents=True, exist_ok=True)

    # data[supervisor][week key][level] = [activity tuple, ...]  (unique, in order)
    # week key: (0, monday ordinal, year, month) for dated rows - a week that
    # runs into the next month is split in two - or (1, sheet index) for a
    # sheet that is one week
    data, labels = {}, {}
    unassigned = set()
    bad_dates = 0
    uploaded = {}                           # uploaded[week key] = {norm(name), ...}

    for sheet_idx, sheet in enumerate(xls.sheet_names):
        df, problem = read_week_sheet(xls, sheet, needed, optional, finder)
        if df is None:
            log(f"  - Skipping sheet '{sheet}':\n      {problem}")
            continue
        dated = by_date and DATE in df.columns

        for _, row in df.iterrows():
            name = clean(row[NAME])
            if not name:
                continue
            supervisor = fixed_name if fixed else clean(row[SUPERVISOR])
            roster.setdefault(norm(name), name)
            entry = tuple(clean(row[f]) if f in df.columns else "" for f in entry_fields)

            week = None
            if any(entry):
                if not dated:
                    week = (1, sheet_idx)
                elif (d := to_date(row[DATE])) is not None:
                    monday = d - timedelta(days=d.weekday())
                    week = (0, monday.toordinal(), d.year, d.month)
                if week:
                    uploaded.setdefault(week, set()).add(norm(name))

            if not supervisor:
                continue
            level = levels.get(norm(name))
            if level is None:
                unassigned.add(name)
                continue
            if not any(entry):
                continue
            if week is None:
                bad_dates += 1
                continue
            labels[week] = date.fromordinal(week[1]) if week[0] == 0 else sheet

            bucket = (data.setdefault(supervisor, {})
                          .setdefault(week, {})
                          .setdefault(level, []))
            if entry not in bucket:          # same activity on several days -> once
                bucket.append(entry)

    if not data:
        raise ReportError("No rows found. Check the sheet layout, the Excel "
                          "columns and the levels file.")

    first_week = profile.get("first_week")
    first_monday = min((w[1] for w in labels if w[0] == 0), default=0)
    for week, monday in labels.items():
        if week[0] == 0:
            if first_week is None:
                number = monday.isocalendar()[1]
            else:
                number = int(first_week) + (week[1] - first_monday) // 7
            labels[week] = week_label(number, monday, week[2], week[3])

    written = []
    for supervisor, weeks in data.items():
        doc = new_document(layout)
        if layout == "weekly":
            for i, week in enumerate(sorted(weeks)):
                # a sheet name gets "Week: " in front; dated weeks already say "Week 39: ..."
                title = labels[week] if week[0] == 0 else f"Week: {labels[week]}"
                add_week_page(doc, supervisor, title, columns, entry_index,
                              weeks[week], first_page=(i == 0))
        else:
            add_title_box(doc, supervisor)
            for week in sorted(weeks):
                add_week(doc, supervisor, labels[week], columns, entry_index, weeks[week])

        path = output_path(out_dir, profile.get("file_name"), supervisor, len(data) > 1)
        try:
            doc.save(path)
        except PermissionError:
            raise ReportError(f"Can't write '{path}' - close it in Word and try again.")
        log(f"  ✓ {path}  ({len(weeks)} week(s))")
        written.append(path)

    if bad_dates:
        log(f"\n  ! {bad_dates} row(s) had an empty or unreadable Date and were left out.")
    if unassigned:
        log("\n  ! These engineers have no level in the levels file and were left out:")
        for n in sorted(unassigned):
            log(f"      - {n}")
        log("    Add them to the levels file (python make_levels.py <workbook>) "
            "and run again.")
    report_missing([(w, labels[w]) for w in sorted(labels)], uploaded, roster, log)
    return written


def inspect_workbook(profile, log=print):
    """Show, sheet by sheet, which Excel header each field matched."""
    fields = active_fields(profile)
    needed = [f["name"] for f in fields if is_required(f)]
    finder = make_finder(fields)
    xls = open_workbook(profile["hours_file"])
    log(f"Checking '{profile['hours_file']}'")
    for sheet in xls.sheet_names:
        raw = read_raw(xls, sheet)
        row, cols, ok = find_header(raw, needed, finder)
        log(f"\n  Sheet '{sheet}':")
        if row is None:
            log("      ! no header row found - first rows look like:")
            for p in preview_rows(raw) or ["(the sheet appears to be empty)"]:
                log(f"          {p}")
            continue
        headers = [clean(h) for h in raw.iloc[row].tolist()]
        log("      " + ("header row found" if ok else
                        "! closest header row - this sheet would be SKIPPED"))
        for f in fields:
            if f["name"] in cols:
                log(f"      ✓ {f['name']:<22} <- '{headers[cols[f['name']]]}'")
            else:
                need = "needed" if is_required(f) else "optional"
                log(f"      ! {f['name']:<22} <- not found ({need})")
        unused = [h for i, h in enumerate(headers) if h and i not in cols.values()]
        if unused:
            log("      Other columns: " + " | ".join(unused))
