"""
Build / update engineer_levels.xlsx from the hours workbook.

Pulls every engineer name out of all the weekly sheets, so the names in the
levels file are spelled exactly as in the hours file. You then only have to
pick a level for each one (the Level column has a Basic / Intermediate /
Senior dropdown).

If the levels file already exists, levels you already filled in are kept and
only new names are added, so it's safe to run again every week.

The window (report_gui.pyw) runs this with its own settings
("Create / update levels file"); this script uses the default ones.

Usage:
    pip install pandas openpyxl
    python make_levels.py hours.xlsx
    python make_levels.py hours.xlsx --levels engineer_levels.xlsx
"""

import argparse
import difflib
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from report_engine import (LEVELS, NAME, SUPERVISOR, ReportError, active_fields,
                           clean, default_profile, make_finder, norm,
                           open_workbook, read_week_sheet)


def collect_names(hours_path, profile, log=print):
    """Return {normalised name: (spelling, [supervisors])} from every sheet."""
    xls = open_workbook(hours_path)
    # Only names and supervisors are needed, so this works on any workbook
    fields = [f for f in active_fields(profile) if f["name"] in (NAME, SUPERVISOR)]
    needed = [f["name"] for f in fields]
    finder = make_finder(fields)
    fixed = profile["supervisor_mode"] == "fixed"
    people = {}
    for sheet in xls.sheet_names:
        df, problem = read_week_sheet(xls, sheet, needed, finder=finder)
        if df is None:
            log(f"  - Skipping sheet '{sheet}':\n      {problem}")
            continue
        for _, row in df.iterrows():
            name = clean(row[NAME])
            if not name:
                continue
            spelling, supervisors = people.setdefault(norm(name), (" ".join(name.split()), []))
            supervisor = profile["supervisor_name"].strip() if fixed else clean(row[SUPERVISOR])
            if supervisor and supervisor not in supervisors:
                supervisors.append(supervisor)
    return people


def read_existing(path):
    """Return {normalised name: (spelling, level)} from an existing levels file."""
    if not path.exists():
        return {}
    try:
        df = pd.read_excel(path)
    except PermissionError:
        raise ReportError(f"Can't open '{path}' - close it in Excel and try again.")
    df.columns = [str(c).strip() for c in df.columns]
    if "Name" not in df.columns or "Level" not in df.columns:
        raise ReportError(f"'{path}' must have the columns 'Name' and 'Level'.")
    existing = {}
    for _, row in df.dropna(subset=["Name"]).iterrows():
        existing[norm(row["Name"])] = (clean(row["Name"]), clean(row["Level"]))
    return existing


def warn_similar(names, log=print):
    """Point out names that look like typos of each other (e.g. Jon / John)."""
    names = sorted(names, key=str.lower)
    lowered = [n.lower() for n in names]
    pairs = []
    for i, a in enumerate(lowered):
        for j in range(i + 1, len(lowered)):
            if difflib.SequenceMatcher(None, a, lowered[j]).ratio() >= 0.85:
                pairs.append((names[i], names[j]))
    if pairs:
        log("\n  ! These names look very similar - check the hours file for typos:")
        for a, b in pairs:
            log(f"      - '{a}'  vs  '{b}'")


def write_levels(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Levels"
    ws.append(["Name", "Level", "Supervisor(s)"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2F3")
    for row in rows:
        ws.append(row)

    # Dropdown so the level can't be misspelled either
    dv = DataValidation(type="list", formula1=f'"{",".join(LEVELS)}"',
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Invalid level",
                        error="Pick Basic, Intermediate or Senior.")
    ws.add_data_validation(dv)
    dv.add(f"B2:B{max(len(rows) + 1, 2) + 200}")   # room for rows added by hand

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 45
    ws.freeze_panes = "A2"

    try:
        wb.save(path)
    except PermissionError:
        raise ReportError(f"Can't write '{path}' - close it in Excel and try again.")


def update_levels(hours_path, levels_path, profile, log=print):
    """Create / update the levels file from the names in the hours workbook."""
    levels_path = Path(levels_path)
    people = collect_names(hours_path, profile, log)
    if not people:
        raise ReportError("No engineer names found. Check the sheet layout.")
    existing = read_existing(levels_path)

    rows, new, missing_level = [], [], []
    for key, (spelling, supervisors) in people.items():
        old_spelling, level = existing.get(key, (None, ""))
        if old_spelling is None:
            new.append(spelling)
        if level.capitalize() in LEVELS:
            level = level.capitalize()
        else:
            missing_level.append(spelling)
        rows.append([spelling, level, ", ".join(supervisors)])

    # Keep people already in the levels file even if they're not in this hours file
    for key, (spelling, level) in existing.items():
        if key not in people:
            rows.append([spelling, level, ""])

    rows.sort(key=lambda r: r[0].lower())
    write_levels(levels_path, rows)

    log(f"  ✓ {levels_path}  ({len(rows)} engineer(s), {len(new)} new)")
    if new and existing:
        log("\n  New names added:")
        for n in sorted(new, key=str.lower):
            log(f"      - {n}")
    if missing_level:
        log(f"\n  ! {len(missing_level)} engineer(s) still need a level - "
            f"open {levels_path.name} and fill in the Level column.")
    warn_similar([s for s, _ in people.values()], log)


def main():
    ap = argparse.ArgumentParser(description="Create/update the engineer levels file.")
    ap.add_argument("hours", help="Excel workbook with one sheet per week")
    ap.add_argument("--levels", default="engineer_levels.xlsx",
                    help="Levels file to create or update (Name | Level)")
    args = ap.parse_args()
    try:
        update_levels(args.hours, args.levels, default_profile("Weekly"))
    except ReportError as e:
        sys.exit(str(e))


if __name__ == "__main__":
    main()
